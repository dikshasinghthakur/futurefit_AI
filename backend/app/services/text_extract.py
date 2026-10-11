import io
import os


class UnsupportedFileType(ValueError):
    pass


class EmptyDocument(ValueError):
    pass


ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}
_MAX_PDF_PAGES = 15


def extract_text(filename: str, data: bytes) -> str:
    ext = os.path.splitext(filename.lower())[1]
    if ext not in ALLOWED_EXTENSIONS:
        raise UnsupportedFileType("Unsupported file type. Please upload a PDF, DOCX or TXT resume.")

    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        text = "\n".join((page.extract_text() or "") for page in reader.pages[:_MAX_PDF_PAGES])
    elif ext == ".docx":
        import docx
        doc = docx.Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        text = "\n".join(parts)
    else:
        text = data.decode("utf-8", errors="replace")

    text = text.replace("\x00", "").strip()
    if len(text) < 40:
        raise EmptyDocument(
            "We couldn't read any text from this file. If it's a scanned image, export a text-based PDF or DOCX instead."
        )
    return text[:60_000]
